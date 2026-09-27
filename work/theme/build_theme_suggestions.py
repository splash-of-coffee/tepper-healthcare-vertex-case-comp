"""Build work/theme/theme_suggestions.pptx: three theme options, two slides each.

Placeholder content only. Colors sampled from the case PDF (see NOTES below).
Run:  python work/theme/build_theme_suggestions.py
Then export the PDF with work/theme/export_pdf.py (PowerPoint COM).
"""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "work" / "theme" / "theme_suggestions.pptx"
LOGO = ROOT / "work" / "theme" / "assets" / "case_cover_logos.png"

# --- Colors -----------------------------------------------------------------
# Sampled from materials/2026 VRTX AMKD Tepper Case Competition.pdf on 2026-09-27:
VERTEX_PURPLE = "3E125F"   # Vertex logo triangle fill (mode of 6,920 core pixels)
CASE_HEADING = "7030A0"    # case PDF heading text color (Office default "Purple")
VERTEX_GRAY = "595A5F"     # Vertex wordmark gray (mean of core pixels)
CLUB_NAVY = "010F3B"       # Tepper Healthcare Club logo navy (mean of core pixels)
CLUB_RED = "F13345"        # Tepper Healthcare Club logo red (mean of core pixels)
# From notes/theme-brief.md (official CMU palette):
CARNEGIE_RED = "C41230"
IRON_GRAY = "6D6E71"
STEEL_GRAY = "E0E0E0"
WEAVER_BLUE = "182C4B"
# Chart de-emphasis gray, validated with the dataviz validator against #3E125F:
ALT_GRAY = "8C8F94"
INK = "262626"
WHITE = "FFFFFF"

TITLE_SERIF = "Georgia"    # stands in for Source Serif 4 (not installed here)
SANS = "Arial"             # stands in for Open Sans (not installed here)

THEMES = [
    {
        "key": "A",
        "label": "Theme option A: sponsor-forward (recommended)",
        "title_font": TITLE_SERIF,
        "body_font": SANS,
        "title_color": VERTEX_PURPLE,
        "body_color": VERTEX_GRAY,
        "accent": VERTEX_PURPLE,
        "number_color": VERTEX_GRAY,
        "cover": "light",
        "notes": (
            "Theme option A, sponsor-forward (recommended in notes/theme-brief.md).\n"
            "Intended fonts: Source Serif 4 Semibold for titles; Open Sans for body and footnotes. "
            "This file renders them in Georgia and Arial because neither intended font is installed on this machine.\n"
            "Colors: Vertex purple #3E125F for titles, the recommendation accent, and headline numbers "
            "(sampled from the Vertex logo on the case PDF cover). Body text #595A5F (sampled Vertex wordmark gray). "
            "Rules and table fills Steel Gray #E0E0E0. Alternatives in charts #8C8F94. "
            "Carnegie Red #C41230 appears on the title slide only, as text.\n"
            "Color notes: the brand reference in the theme brief, #52247F, is lighter than the sampled logo purple #3E125F. "
            "The case PDF headings use #7030A0, which is Microsoft Office's default purple, not a Vertex brand color."
        ),
    },
    {
        "key": "B",
        "label": "Theme option B: Tepper-forward co-brand",
        "title_font": SANS,
        "body_font": SANS,
        "title_color": CARNEGIE_RED,
        "body_color": IRON_GRAY,
        "accent": VERTEX_PURPLE,
        "number_color": CARNEGIE_RED,
        "cover": "red",
        "notes": (
            "Theme option B, Tepper-forward co-brand.\n"
            "Intended font: Open Sans only (bold titles, regular body). This file renders it in Arial because Open Sans is not installed on this machine.\n"
            "Colors: Carnegie Red #C41230 for the title slide, titles, and slide numbers. "
            "Vertex purple #3E125F marks the recommendation in charts. Body text Iron Gray #6D6E71. Fills Steel Gray #E0E0E0. "
            "Alternatives in charts #8C8F94.\n"
            "Risk from the theme brief: readers read red as a loss in financial charts, so red stays out of every chart."
        ),
    },
    {
        "key": "C",
        "label": "Theme option C: Healthcare Club navy",
        "title_font": SANS,
        "body_font": SANS,
        "title_color": WEAVER_BLUE,
        "body_color": "4A4A4A",
        "accent": VERTEX_PURPLE,
        "number_color": WEAVER_BLUE,
        "cover": "navy",
        "notes": (
            "Theme option C, Healthcare Club navy.\n"
            "Intended font: Open Sans only. This file renders it in Arial because Open Sans is not installed on this machine.\n"
            "Colors: title slide in the Healthcare Club navy #010F3B (sampled from the club logo on the case PDF cover). "
            "Titles in CMU Weaver Blue #182C4B. Vertex purple #3E125F marks the recommendation in charts. "
            "Body text #4A4A4A. Alternatives in charts #8C8F94. "
            "Red appears on the title slide only, in the club red #F13345 sampled from the club logo."
        ),
    },
]


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def text_box(slide, x, y, w, h, text, font, size, color, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False, spacing_after=0):
    """Add a text box. `text` may be a string or a list of strings (one paragraph each)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    paras = text if isinstance(text, list) else [text]
    for i, para_text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing_after)
        run = p.add_run()
        run.text = para_text
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
    return tb


def bullets(slide, x, y, w, h, items, font, size, color):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        # Real bullet through paragraph properties (no literal bullet characters).
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.22)))
        pPr.set("indent", str(-Inches(0.22)))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        r1 = p.add_run()
        r1.text = lead
        r1.font.bold = True
        r1.font.name = font
        r1.font.size = Pt(size)
        r1.font.color.rgb = rgb(color)
        r2 = p.add_run()
        r2.text = rest
        r2.font.name = font
        r2.font.size = Pt(size)
        r2.font.color.rgb = rgb(color)
    return tb


def rect(slide, x, y, w, h, fill, rounded=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.shadow.inherit = False
    if rounded:
        shp.adjustments[0] = 0.08
    return shp


def title_slide(prs, t):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    W, H = prs.slide_width, prs.slide_height
    dark = t["cover"] in ("red", "navy")
    if dark:
        bg = CARNEGIE_RED if t["cover"] == "red" else CLUB_NAVY
        rect(s, 0, 0, W, H, bg)
    title_c = WHITE if dark else t["title_color"]
    sub_c = "F2F2F2" if dark else VERTEX_GRAY
    label_c = WHITE if dark else CARNEGIE_RED
    if t["cover"] == "navy":
        label_c = CLUB_RED

    text_box(s, Inches(0.8), Inches(0.75), Inches(9), Inches(0.35),
             "2026 TEPPER HEALTHCARE CASE COMPETITION  ·  QUALIFYING ROUND",
             t["body_font"], 12, label_c, bold=True)
    text_box(s, Inches(0.8), Inches(1.45), Inches(10.5), Inches(1.9),
             ["Reaching AMKD Patients", "in Primary Care"],
             t["title_font"], 44, title_c, bold=True, anchor=MSO_ANCHOR.TOP)
    text_box(s, Inches(0.8), Inches(3.05), Inches(10.5), Inches(0.5),
             "A go-to-market recommendation for inaxaplin, prepared for Vertex Pharmaceuticals",
             t["body_font"], 18, sub_c)
    text_box(s, Inches(0.8), Inches(3.75), Inches(10.5), Inches(0.4),
             "Vidhur Vashisht  ·  Skylar Dennerlein  ·  Mike Homze  ·  Christian Woodfin  ·  October 2026",
             t["body_font"], 14, sub_c)

    # Logo panel: the case cover image (Vertex | Healthcare Club), unaltered.
    logo_w = Inches(5.2)
    logo_h = Emu(int(logo_w * 253 / 1277))
    if dark:
        card = rect(s, Inches(0.8), Inches(5.35), logo_w + Inches(0.5), logo_h + Inches(0.5), WHITE, rounded=True)
        s.shapes.add_picture(str(LOGO), Inches(1.05), Inches(5.6), width=logo_w)
    else:
        s.shapes.add_picture(str(LOGO), Inches(0.8), Inches(5.6), width=logo_w)
    cmu_c = WHITE if dark else CARNEGIE_RED
    text_box(s, Inches(8.3), Inches(5.75), Inches(4.2), Inches(0.9),
             ["Carnegie Mellon University", "Tepper School of Business"],
             t["title_font"] if t["key"] == "A" else t["body_font"], 16, cmu_c, bold=True, align=PP_ALIGN.RIGHT)
    text_box(s, Inches(0.8), Inches(7.05), Inches(8), Inches(0.3), t["label"],
             t["body_font"], 10, "F2F2F2" if dark else IRON_GRAY, italic=True)
    s.notes_slide.notes_text_frame.text = t["notes"]
    return s


def body_slide(prs, t, number):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    W = prs.slide_width
    text_box(s, Inches(0.6), Inches(0.45), Inches(12.1), Inches(1.1),
             "Placeholder title: the highlighted combination returns $62M in this sample, $17M more than the next one",
             t["title_font"], 26, t["title_color"], bold=True)

    bullets(s, Inches(0.6), Inches(1.95), Inches(4.6), Inches(3.6), [
        ("Titles: ", "one full sentence, at most two lines"),
        ("Bullets: ", "fragments of 12 words or fewer, each with a number"),
        ("Accent: ", "Vertex purple marks our recommendation on every chart"),
        ("Alternatives: ", "gray #8C8F94, never red"),
        ("Assumptions: ", "each team value carries its register ID (A-01)"),
    ], t["body_font"], 14, t["body_color"])

    # Headline number callout
    text_box(s, Inches(0.6), Inches(4.75), Inches(4.6), Inches(0.7), "$62M",
             t["body_font"], 40, t["accent"], bold=True)
    text_box(s, Inches(0.6), Inches(5.45), Inches(4.6), Inches(0.4),
             "sample headline figure, placeholder data (A-01)", t["body_font"], 12, t["body_color"])

    # Chart title as text (sentence case), then a native bar chart.
    text_box(s, Inches(5.7), Inches(1.95), Inches(7.0), Inches(0.35),
             "Added risk-adjusted NPV by combination, $M (placeholder data)",
             t["body_font"], 13, INK, bold=True)
    cd = CategoryChartData()
    cats = ["Combination E", "Combination D", "Combination C", "Combination B", "Recommended (example)"]
    vals = [9, 18, 31, 45, 62]
    cd.categories = cats
    cd.add_series("NPV", vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(5.6), Inches(2.35), Inches(7.1), Inches(3.55), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 110
    plot.vary_by_categories = False
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(ALT_GRAY)
    pt = ser.points[len(vals) - 1]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(t["accent"])
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '"$"#,##0"M"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(12)
    dl.font.name = t["body_font"]
    dl.font.color.rgb = rgb(INK)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = 75
    va.minimum_scale = 0
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.name = t["body_font"]
    ca.tick_labels.font.color.rgb = rgb(INK)
    ca.format.line.color.rgb = rgb("BFBFBF")
    ca.has_major_gridlines = False

    # Footnote and slide number.
    text_box(s, Inches(0.6), Inches(6.55), Inches(10.8), Inches(0.55), [
        "Sources: [1] Sample source name, 4/19/26, p. 5. A-01 = team assumption ID from the assumptions register.",
        "Placeholder data for theme review only; these are not results.  " + t["label"] + ".",
    ], t["body_font"], 9, IRON_GRAY)
    text_box(s, Inches(11.9), Inches(6.85), Inches(0.8), Inches(0.3), str(number),
             t["body_font"], 11, t["number_color"], bold=True, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = t["notes"]
    return s


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    n = 0
    for t in THEMES:
        title_slide(prs, t)
        n += 1
        n += 1
        body_slide(prs, t, n)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
