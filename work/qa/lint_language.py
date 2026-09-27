"""Language check for the Vertex / AMKD case-competition project.

Implements the enforcement section of notes/writing-standard.md:
- scans every .md file under work/ and the text of every .pptx under work/;
- reports each banned word or sentence pattern with its file and line;
- skips text inside double quotation marks, inline code, URLs, and fenced code;
- skips the allowed technical terms (pivotal trial, key opinion leader, and so on);
- applies the slide exemption: PPTX table cells and chart labels are skipped, and
  slide bullets may be fragments of 12 words or fewer.

Severity:
- "ban" hits count toward the zero-hit gate.
- "review" items are listed for a human check and do not count.

A hit is kept with a written reason by adding "lint-ok: <reason>" to the same
line (in Markdown, inside an HTML comment: <!-- lint-ok: reason -->). Kept hits
are listed separately with their reasons.

Usage:
    python work/qa/lint_language.py                 # scan work/
    python work/qa/lint_language.py path1 path2     # scan given files or folders
    python work/qa/lint_language.py --acronyms deck.pptx   # also list acronyms
                                                           # missing from notes/acronyms.md
Exit code: 0 when there are zero ban hits, 1 otherwise.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "work"
ACRONYMS_FILE = ROOT / "notes" / "acronyms.md"

# ---------------------------------------------------------------------------
# Rules. Each rule is (name, regex, severity, note). Regexes are case-insensitive
# unless the pattern sets its own flags.
# ---------------------------------------------------------------------------

BAN = "ban"
REVIEW = "review"

HYPE = [
    "robust", "compelling", "powerful", "seamless", "seamlessly", "cutting-edge",
    "innovative", "groundbreaking", "game-changing", "game changer", "world-class",
    "best-in-class", "holistic", "dynamic", "unparalleled",
]
IMPORTANCE = ["crucial", "vital", "essential"]
VERBS = [
    r"delv(?:e|es|ed|ing)", r"unlock(?:s|ed|ing)?", r"harness(?:es|ed|ing)?",
    r"empower(?:s|ed|ing|ment)?", r"elevat(?:e|es|ed|ing)", r"streamlin(?:e|es|ed|ing)",
    r"spearhead(?:s|ed|ing)?", r"bolster(?:s|ed|ing)?", r"underscor(?:e|es|ed|ing)",
    r"showcas(?:e|es|ed|ing)", r"embark(?:s|ed|ing)?", r"reimagin(?:e|es|ed|ing)",
    r"supercharg(?:e|es|ed|ing)", r"revolutioniz(?:e|es|ed|ing)", r"navigat(?:e|es|ed|ing)",
]
NOUNS = [
    "landscape", "landscapes", "ecosystem", "ecosystems", "tapestry", "journey", "journeys",
    "realm", "synergy", "synergies", "paradigm", "testament", "arsenal", "north star",
    "deep dive", "deep-dive", "game plan", "secret sauce", "needle-mover", "needle mover",
    "cornerstone",
]
INTENSIFIERS = [
    "genuinely", "truly", "really", "actually", "incredibly", "extremely", "highly",
    "deeply", "notably", "importantly", "clearly", "simply", "just",
]
FRAMING = [
    r"it'?s worth noting", r"it is worth noting", r"worth noting", r"at the end of the day",
    r"in today'?s fast-paced world", r"let'?s dive in", r"here'?s the thing",
    r"the bottom line is", r"moving forward", r"at its core", r"when it comes to",
    r"in order to",
]
SEQUENCING = [
    r"step by step", r"step-by-step", r"piece by piece", r"one at a time",
    r"here'?s how it breaks down", r"here is how it breaks down", r"let'?s break (?:it|this) down",
]

NUMBER_WORDS = r"(?:one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve)"


def word(w: str) -> str:
    return r"\b" + w + r"\b"


RULES: list[tuple[str, re.Pattern, str, str]] = []


def add(name: str, pattern: str, severity: str, note: str, flags: int = re.IGNORECASE) -> None:
    RULES.append((name, re.compile(pattern, flags), severity, note))


for w in HYPE:
    add(w, word(re.escape(w)), BAN, "hype adjective: use the number or the property")
for w in IMPORTANCE:
    add(w, word(w), BAN, "importance word: say why it matters")
add("critical", r"\bcritical\b(?!\s+(?:path|care|access hospital))", BAN, "importance word: say why it matters")
add("key (adjective)", r"\bkey\b(?!\s+(?:account manager|account managers|opinion leader|opinion leaders|performance indicator|performance indicators))",
    BAN, "importance word; allowed only in KAM, KOL, KPI")
add("pivotal", r"\bpivotal\b(?!\s+(?:[\w-]+\s+)?(?:trial|trials|study|studies|stage|development|program|programs|data|readout|phase|results|population|endpoint))",
    BAN, "allowed only as 'pivotal trial' and similar")
add("comprehensive", r"\bcomprehensive\b(?!\s+metabolic panel)", BAN, "hype adjective")
for w in VERBS:
    add(re.sub(r"\(\?:[^)]*\)\??", "", w), word(w), BAN, "AI-accent verb: use a plain verb")
add("leverage", r"\bleverag(?:e|es|ed|ing)\b", BAN, "AI-accent verb (financial leverage is allowed: add lint-ok)")
add("foster", r"(?<!Logan )\bfoster(?:s|ed|ing)?\b", BAN, "AI-accent verb")
add("optimize", r"\boptimi[sz](?:e|es|ed|ing|ation)\b", REVIEW, "allowed only for mathematical optimization")
for w in NOUNS:
    add(w, word(re.escape(w)), BAN, "AI-accent noun: use the plain word")
add("lever (metaphor)", r"\blevers?\b", REVIEW, "banned as a metaphor; name the driver")
add("unique", r"\buniquely?\b", REVIEW, "allowed only if literally one of a kind")
for w in INTENSIFIERS:
    add(w, word(w), BAN, "intensifier: delete")
add("significantly", r"(?<!statistically )\bsignificantly\b", BAN, "intensifier unless statistical")
for p in FRAMING:
    add(p, r"\b" + p + r"\b", BAN, "framing phrase: state the content")
for p in SEQUENCING:
    add(p, r"\b" + p + r"\b", BAN, "cinematic sequencing (global rule)")
add("not just/not only", r"\bnot (?:just|only)\b", BAN, "'Not just X, but Y' pattern")
add("it's not X, it's Y", r"\bit'?s not\b[^.;]{0,80},\s*it'?s\b", BAN, "'It's not X, it's Y' pattern")
add("Not X. fragment", r"(?:^|[.!?]\s+)Not\s+[^.!?]{1,40}\.", BAN, "'Not X. Y.' fragment", flags=0)
add("colon reveal", r"\b[Tt]he (?:result|answer|bottom line|takeaway|catch|kicker|twist|verdict|upshot|truth|lesson|point|secret|reality)\s*:\s",
    BAN, "colon reveal: write one declarative sentence", flags=0)
NUM_CI = "(?i:" + NUMBER_WORDS[3:]  # case-insensitive copy of the number-word group
add("dramatic counting",
    r"(?:^|[.!?]\s+|\*\*)" + NUM_CI + r"\s+[A-Za-z-]+,\s+" + NUM_CI + r"\s+[A-Za-z-]+",
    BAN, "'One X, two Y' counting (global rule)", flags=0)
add("question", r"\?(?=\s|$)", REVIEW, "rhetorical questions are banned; real questions to the team are fine")

# ---------------------------------------------------------------------------
# Text cleaning: remove the parts of a line the standard exempts.
# ---------------------------------------------------------------------------

QUOTE_PATTERNS = [
    re.compile(r"\"[^\"]*\""),          # straight double quotes
    re.compile(r"“[^”]*”"),  # curly double quotes
    re.compile(r"`[^`]*`"),             # inline code
    re.compile(r"\]\([^)]*\)"),          # markdown link targets
    re.compile(r"https?://\S+"),        # bare URLs
    re.compile(r"<!--.*?-->"),          # HTML comments
]


def clean(line: str) -> str:
    out = line
    for pat in QUOTE_PATTERNS:
        out = pat.sub(lambda m: " " * len(m.group(0)), out)
    return out


@dataclass
class Hit:
    location: str
    rule: str
    severity: str
    note: str
    context: str
    kept_reason: str | None = None


def scan_text(location: str, text: str, is_slide_fragment: bool = False, is_title: bool = False) -> list[Hit]:
    hits: list[Hit] = []
    kept = None
    m = re.search(r"lint-ok:\s*([^>]*?)(?:-->|$)", text)
    if m:
        kept = m.group(1).strip() or "no reason given"
    cleaned = clean(text)
    for name, pat, severity, note in RULES:
        if is_slide_fragment and name in ("Not X. fragment",):
            continue
        for mm in pat.finditer(cleaned):
            sev = severity
            if name == "question" and is_title:
                sev = BAN  # slide titles must be declarative sentences
            start = max(0, mm.start() - 40)
            ctx = text[start: mm.end() + 40].replace("\n", " ").strip()
            hits.append(Hit(location, name, sev, note, ctx, kept))
    return hits


# ---------------------------------------------------------------------------
# Markdown scanning
# ---------------------------------------------------------------------------

def scan_markdown(path: Path) -> list[Hit]:
    hits: list[Hit] = []
    in_fence = False
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except UnicodeDecodeError:
        lines = path.read_text(encoding="latin-1").splitlines()
    rel = path.relative_to(ROOT) if path.is_relative_to(ROOT) else path
    exempt_questions = any(k in path.name.lower() for k in ("qa_bank", "interview_guide", "question"))
    for i, line in enumerate(lines, 1):
        if line.strip().startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        for h in scan_text(f"{rel}:{i}", line):
            if h.rule == "question" and exempt_questions:
                continue
            hits.append(h)
    return hits


# ---------------------------------------------------------------------------
# PPTX scanning (slide exemption applied)
# ---------------------------------------------------------------------------

def scan_pptx(path: Path) -> list[Hit]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import PP_PLACEHOLDER
    except ImportError:
        print("python-pptx is not installed; skipping", path, file=sys.stderr)
        return []
    hits: list[Hit] = []
    rel = path.relative_to(ROOT) if path.is_relative_to(ROOT) else path
    prs = Presentation(str(path))
    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    for s_idx, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            for h in _scan_shape(shape, f"{rel}#slide{s_idx}", title_types):
                hits.append(h)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            for j, line in enumerate(notes.splitlines(), 1):
                hits.extend(scan_text(f"{rel}#slide{s_idx}:notes:{j}", line))
    return hits


def _scan_shape(shape, loc: str, title_types) -> list[Hit]:
    hits: list[Hit] = []
    # Groups: recurse.
    if shape.shape_type is not None and getattr(shape, "shapes", None) is not None and shape.shape_type == 6:
        for sub in shape.shapes:
            hits.extend(_scan_shape(sub, loc, title_types))
        return hits
    if getattr(shape, "has_table", False) and shape.has_table:
        return hits  # slide exemption: table cells skipped
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return hits  # slide exemption: chart labels skipped
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return hits
    is_title = False
    if shape.is_placeholder:
        try:
            is_title = shape.placeholder_format.type in title_types
        except Exception:
            is_title = False
    if not is_title and shape.name.lower().startswith("title"):
        is_title = True
    for p_idx, para in enumerate(shape.text_frame.paragraphs, 1):
        text = "".join(run.text for run in para.runs).strip()
        if not text:
            continue
        where = f"{loc}:{shape.name}:{p_idx}"
        words = text.split()
        if is_title:
            hits.extend(scan_text(where, text, is_title=True))
            if len(words) < 5:
                hits.append(Hit(where, "short title", REVIEW, "titles should be full sentences", text))
            if len(words) > 25:
                hits.append(Hit(where, "long title", REVIEW, "titles should fit in two lines", text))
        else:
            fragment = len(words) <= 12 and not text.endswith(".")
            hits.extend(scan_text(where, text, is_slide_fragment=fragment))
            if len(words) > 12 and not re.search(r"[.!?]$", text):
                hits.append(Hit(where, "long fragment", REVIEW, "fragments must be 12 words or fewer; otherwise write a sentence", text))
            if fragment and not _carries_number_or_name(text):
                hits.append(Hit(where, "empty fragment", REVIEW, "a fragment must carry a number or a named thing", text))
    return hits


def _carries_number_or_name(text: str) -> bool:
    """True when a slide fragment holds a digit, a number word, an acronym, or a
    capitalized name after its first word (the slide exemption's requirement)."""
    if re.search(r"\d", text):
        return True
    if re.search(r"\b" + NUMBER_WORDS + r"\b", text, re.IGNORECASE):
        return True
    words = re.findall(r"[A-Za-z][A-Za-z0-9'-]*", text)
    if any(re.fullmatch(r"[A-Z][A-Z0-9-]+s?", w) for w in words):
        return True
    return any(w[0].isupper() for w in words[1:])


# ---------------------------------------------------------------------------
# Acronym check (used in the final language check)
# ---------------------------------------------------------------------------

def known_acronyms() -> set[str]:
    if not ACRONYMS_FILE.exists():
        return set()
    text = ACRONYMS_FILE.read_text(encoding="utf-8")
    found = set(re.findall(r"\*\*([^*]+)\*\*", text))
    out: set[str] = set()
    for term in found:
        for part in re.split(r"[ /,()]+", term):
            if part:
                out.add(part.strip())
    return out


def acronyms_in_pptx(path: Path) -> set[str]:
    from pptx import Presentation
    prs = Presentation(str(path))
    tokens: set[str] = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                tokens.update(re.findall(r"\b[A-Z][A-Z0-9]{1,}[a-z]?\b", shape.text_frame.text))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        tokens.update(re.findall(r"\b[A-Z][A-Z0-9]{1,}[a-z]?\b", cell.text))
    return tokens


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def collect(paths: list[str]) -> list[Path]:
    targets: list[Path] = []
    if not paths:
        paths = [str(WORK)]
    for p in paths:
        pp = Path(p)
        if not pp.is_absolute():
            pp = (Path.cwd() / pp).resolve()
        if pp.is_dir():
            targets.extend(sorted(pp.rglob("*.md")))
            targets.extend(sorted(x for x in pp.rglob("*.pptx") if not x.name.startswith("~$")))
        elif pp.exists():
            targets.append(pp)
    return targets


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="*")
    ap.add_argument("--acronyms", action="store_true", help="list acronyms in PPTX files missing from notes/acronyms.md")
    ap.add_argument("--quiet", action="store_true", help="print only the totals")
    args = ap.parse_args()

    hits: list[Hit] = []
    files = collect(args.paths)
    for f in files:
        if f.suffix.lower() == ".md":
            hits.extend(scan_markdown(f))
        elif f.suffix.lower() == ".pptx":
            hits.extend(scan_pptx(f))

    bans = [h for h in hits if h.severity == BAN and not h.kept_reason]
    kept = [h for h in hits if h.kept_reason]
    reviews = [h for h in hits if h.severity == REVIEW and not h.kept_reason]

    if not args.quiet:
        if bans:
            print("== BAN HITS (count toward the zero-hit gate) ==")
            for h in bans:
                print(f"{h.location}: [{h.rule}] {h.note} :: {h.context}")
        if kept:
            print("\n== HITS KEPT WITH A WRITTEN REASON ==")
            for h in kept:
                print(f"{h.location}: [{h.rule}] reason: {h.kept_reason} :: {h.context}")
        if reviews:
            print("\n== REVIEW ITEMS (do not count) ==")
            for h in reviews:
                print(f"{h.location}: [{h.rule}] {h.note} :: {h.context}")

    if args.acronyms:
        known = known_acronyms()
        print("\n== ACRONYMS NOT IN notes/acronyms.md ==")
        for f in files:
            if f.suffix.lower() == ".pptx":
                missing = sorted(t for t in acronyms_in_pptx(f) if t not in known)
                print(f"{f.name}: {', '.join(missing) if missing else 'none'}")

    print(f"\nFiles scanned: {len(files)}. Ban hits: {len(bans)}. Kept with reason: {len(kept)}. Review items: {len(reviews)}.")
    return 0 if not bans else 1


if __name__ == "__main__":
    sys.exit(main())
